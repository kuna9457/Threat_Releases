"""Bulk CAN-wise URL Risk Assessment.
Reads CAN_Merged.xlsx (CAN ID | URL | App name | Request id), scans every CAN's URLs with the
existing Threat_Backend scanner and writes one ICICI URL Risk Assessment PDF per CAN ID:
    Reports\\<AppName>_<CANID>_ThreatScanner.pdf
"""
import os, re, sys, time, warnings, zipfile
import xml.etree.ElementTree as ET
from datetime import datetime

warnings.filterwarnings("ignore")
try:
    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
except Exception:
    pass
try:  # Zscaler / proxy SSL fix - skip certificate verification for outbound API calls
    from requests import adapters as _ra
    _orig_send = _ra.HTTPAdapter.send

    def _send(self, req, **kw):
        kw["verify"] = False
        return _orig_send(self, req, **kw)
    _ra.HTTPAdapter.send = _send
except Exception as e:
    print(f"[WARN] Could not patch requests: {e}")

BASE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.join(BASE, "Threat_Backend")
sys.path.insert(0, BACKEND)
os.chdir(BACKEND)
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(BACKEND, ".env"))
except Exception:
    pass

from app.services.scanner import scan_urls_batch          # noqa: E402
from app.services.pdf_report import generate_pdf_report   # noqa: E402
from app.cache.memory_cache import cache as scan_cache     # noqa: E402
from app.utils.resilient import pause, wait_for_internet, is_transient_error   # noqa: E402

XLSX = os.path.join(BASE, "CAN_Merged.xlsx")
OUT_DIR = os.path.join(BASE, "Reports")
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def read_xlsx(path):
    """Return list of rows (dict col_letter -> str) from the first sheet. Stdlib only."""
    with zipfile.ZipFile(path) as z:
        shared = []
        if "xl/sharedStrings.xml" in z.namelist():
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in root.iter(NS + "si"):
                shared.append("".join(t.text or "" for t in si.iter(NS + "t")))
        sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        rows = []
        for row in sheet.iter(NS + "row"):
            vals = {}
            for c in row.findall(NS + "c"):
                col = re.sub(r"\d", "", c.get("r"))
                t = c.get("t")
                v = c.find(NS + "v")
                if t == "s" and v is not None:
                    val = shared[int(v.text)]
                elif t == "inlineStr":
                    val = "".join(x.text or "" for x in c.iter(NS + "t"))
                else:
                    val = v.text if v is not None else ""
                vals[col] = (val or "").strip()
            rows.append(vals)
        return rows


def clean_urls(cell):
    out, seen = [], set()
    for tok in re.split(r"[,;\s]+", cell or ""):
        t = tok.strip()
        if not t or re.fullmatch(r"[\.\*/]+", t):
            continue
        # wildcard entries (*.domain.com / .domain.com) are assessed on their base domain
        t = re.sub(r"^(https?://)?(\*\.|\.)+", r"\1", t, flags=re.I).replace("*", "www")
        if not t or t.lower() in seen:
            continue
        seen.add(t.lower())
        out.append(t)
    return out


SOURCES = ("VirusTotal", "AbuseIPDB", "SSLyze", "WHOIS")
SUMMARY_HDR = ["CAN ID", "App Name", "Request ID", "URLs", "High", "Medium", "Low", "No Risk", "Unknown",
               "Worst Risk", "URLs with unassessed sources", "Report / Note"]
DETAIL_HDR = ["CAN ID", "App Name", "URL", "Final Risk Level", *SOURCES]
NF_HDR = ["CAN ID", "App Name", "URL", "Source", "Issue", "Detail (excluded from risk calculation)"]


def classify_issue(detail):
    t = (detail or "").lower()
    if "resolve" in t:
        return "Unable to resolve"
    if "not found" in t or "no creation date" in t or "no result" in t:
        return "Not found"
    if "connect" in t or "timed out" in t or "timeout" in t:
        return "Unable to connect"
    if "bare-ip" in t or "not applicable" in t:
        return "Not applicable"
    return "No usable result"


def add_excel_rows(rec, results, worst, fpath, sum_rows, det_rows, nf_rows):
    counts = {k: 0 for k in ("High", "Medium", "Low", "No Risk", "Unknown")}
    flagged = 0
    for res in results:
        lvl = res.get("final_risk_level", "Unknown")
        counts[lvl if lvl in counts else "Unknown"] += 1
        rs = res.get("risk_sources", {}) or {}
        row = [rec["can"], rec["app"], res.get("url", ""), lvl]
        bad_here = False
        for s in SOURCES:
            info = rs.get(s, {}) or {}
            if info.get("usable") and info.get("level"):
                row.append(f"{info['level']} ({info.get('detail', '')})")
            else:
                row.append("Not assessed")
                bad_here = True
                nf_rows.append([rec["can"], rec["app"], res.get("url", ""), s,
                                classify_issue(info.get("detail")), str(info.get("detail", ""))[:200]])
        flagged += bad_here
        det_rows.append(row)
    sum_rows.append([rec["can"], rec["app"], rec["req"], len(results), counts["High"], counts["Medium"],
                     counts["Low"], counts["No Risk"], counts["Unknown"], worst, flagged, fpath])


def write_excel(path, sum_rows, det_rows, nf_rows):
    """Excel summary (3 sheets), rewritten after every CAN so progress is never lost.
    Falls back to CSV files if openpyxl is not installed."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        import csv
        for name, hdr, rows in (("Summary", SUMMARY_HDR, sum_rows), ("URL_Details", DETAIL_HDR, det_rows),
                                ("NotFound_Unresolved", NF_HDR, nf_rows)):
            with open(path.replace(".xlsx", f"_{name}.csv"), "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(hdr)
                w.writerows(rows)
        return
    wb = openpyxl.Workbook()
    widths = {"Summary": [9, 34, 40, 7, 7, 8, 6, 9, 9, 12, 14, 70],
              "URL Details": [9, 30, 60, 14, 26, 26, 30, 22],
              "Not Found - Unresolved": [9, 30, 60, 13, 18, 70]}
    first = True
    for title, hdr, rows in (("Summary", SUMMARY_HDR, sum_rows), ("URL Details", DETAIL_HDR, det_rows),
                             ("Not Found - Unresolved", NF_HDR, nf_rows)):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        ws.append(hdr)
        for row in rows:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, w in enumerate(widths[title], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(path)
    except PermissionError:
        print(f"[NOTE] Close the Excel file to let the summary update: {os.path.basename(path)}")


def result_problems(res):
    """Sources in one scan result that failed for a temporary reason (limit / timeout / connection)."""
    d = res.get("data", {}) or {}
    bad = []
    # API sources only. SSLyze/WHOIS errors (host has no TLS, no WHOIS record...) are real answers
    # after their own internal retries, so they must not block a report forever.
    for src in ("virustotal", "abuseipdb"):
        err = (d.get(src) or {}).get("error")
        if err and (src == "virustotal" or is_transient_error(err)):
            bad.append(f"{res.get('url')}: {src}: {str(err)[:80]}")
    return bad


def scan_complete(urls, max_rounds=8):
    """Scan all URLs; any URL whose lookups failed temporarily is re-scanned (after a pause)
    until every source returned a real answer. Returns (results, still_bad_messages)."""
    results = {}
    todo = list(urls)
    bad = []
    for rnd in range(1, max_rounds + 1):
        try:
            batch = scan_urls_batch(todo)
        except Exception as e:
            print(f"[PAUSED] scan error: {e}")
            bad = [f"scan error: {e}"]
            pause("Scanner", "unexpected error", rnd, base=20)
            continue
        bad, retry = [], []
        for u, res in zip(todo, batch):
            results[u] = res
            p = result_problems(res)
            if p:
                bad += p
                retry.append(u)
                scan_cache.pop(u, None)      # never re-use a result with failed sources
        if not retry:
            return [results[u] for u in urls], []
        print(f"[RETRY] {len(retry)} URL(s) had failed lookups; waiting, then re-scanning (round {rnd}/{max_rounds})")
        wait_for_internet("Scanner")
        pause("Scanner", f"{len(retry)} URL(s) incomplete", rnd, base=30, cap=300)
        todo = retry
    if len(results) < len(urls) and not bad:
        bad = ["some URLs were not scanned"]
    return [results[u] for u in urls if u in results], bad


def main():
    print("=" * 60)
    print("  ICICI BANK | THREAT SCANNER | CAN-WISE REPORT GENERATOR")
    print("=" * 60)
    if not os.path.exists(XLSX):
        print(f"[ERROR] CAN_Merged.xlsx not found: {XLSX}")
        return 1
    rows = read_xlsx(XLSX)
    hdr = {v.strip().lower(): k for k, v in rows[0].items()}
    for need in ("can id", "url", "app name", "request id"):
        if need not in hdr:
            print(f"[ERROR] Column '{need}' missing in Excel header.")
            return 1
    recs, skipped = [], 0
    for r in rows[1:]:
        def g(n):
            return r.get(hdr[n], "").strip()
        can, app, req, urls = g("can id"), g("app name"), g("request id"), g("url")
        if not can:
            continue
        if not app or app.upper() == "NOT FOUND" or not req or req.upper() == "NOT FOUND" or not urls:
            skipped += 1
            continue
        recs.append({"can": can, "app": app, "req": req, "urls": clean_urls(urls)})
    print(f"Loaded {len(recs)} CAN ID(s); {skipped} skipped (NOT FOUND / empty).")

    only = input("Enter a single CAN ID to run only that one (or press Enter for ALL): ").strip()
    if only:
        recs = [r for r in recs if r["can"] == only]
        if not recs:
            print(f"CAN {only} not found / skipped.")
            return 1
    skip_ex = input("Skip CAN IDs whose PDF already exists? (Y/n): ").strip().lower() != "n"

    os.makedirs(OUT_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log = os.path.join(OUT_DIR, f"ThreatScanner_{stamp}.log")
    summary = []
    sum_rows, det_rows, nf_rows = [], [], []
    xlsx_path = os.path.join(OUT_DIR, f"ThreatScanner_Summary_{stamp}.xlsx")
    total = len(recs)
    for i, r in enumerate(recs, 1):
        safe_app = re.sub(r"\s+", "_", re.sub(r'[\\/:*?"<>|]', "", r["app"]))
        fpath = os.path.join(OUT_DIR, f"{safe_app}_{r['can']}_ThreatScanner.pdf")
        if skip_ex and os.path.exists(fpath):
            print(f"[{i}/{total}] CAN {r['can']} | {r['app']} | already exists - skipped")
            continue
        t0 = time.time()
        print(f"[{i}/{total}] CAN {r['can']} | {r['app']} | {len(r['urls'])} URL(s) - scanning...")
        try:
            results, bad = scan_complete(r["urls"])
            if bad:
                msg = "INCOMPLETE - PDF not saved (will be retried on next run): " + "; ".join(bad[:3])
                print(f"           {msg}")
                summary.append((r["can"], r["app"], r["req"], len(r["urls"]), "INCOMPLETE", msg))
                sum_rows.append([r["can"], r["app"], r["req"], len(r["urls"]), "", "", "", "", "", "INCOMPLETE", "", msg])
                with open(log, "a", encoding="utf-8") as lf:
                    lf.write(f"{datetime.now():%H:%M:%S} | {summary[-1]}\n")
                write_excel(xlsx_path, sum_rows, det_rows, nf_rows)
                continue
            pdf = generate_pdf_report(results, final_comment="", app_name=r["app"],
                                      can_id=r["can"], server_ip=None, request_id=r["req"])
            with open(fpath, "wb") as f:
                f.write(pdf)
            levels = [x.get("final_risk_level", "Unknown") for x in results]
            worst = next((l for l in ("High", "Medium", "Low", "No Risk") if l in levels), "Unknown")
            print(f"           done in {time.time() - t0:.0f}s | worst risk: {worst} | {os.path.basename(fpath)}")
            summary.append((r["can"], r["app"], r["req"], len(r["urls"]), worst, fpath))
            add_excel_rows(r, results, worst, fpath, sum_rows, det_rows, nf_rows)
        except Exception as e:
            print(f"           ERROR: {e}")
            summary.append((r["can"], r["app"], r["req"], len(r["urls"]), "ERROR", str(e)))
            sum_rows.append([r["can"], r["app"], r["req"], len(r["urls"]), "", "", "", "", "", "ERROR", "", str(e)])
        with open(log, "a", encoding="utf-8") as lf:
            lf.write(f"{datetime.now():%H:%M:%S} | {summary[-1]}\n")
        write_excel(xlsx_path, sum_rows, det_rows, nf_rows)

    if sum_rows:
        print(f"\nExcel summary: {xlsx_path}")
    print(f"Reports folder: {OUT_DIR}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
